"""文档加载器：支持 .pdf / .docx / .txt / .md 四种格式。

每个返回的 Document.metadata 中保留溯源信息：
- source    文件绝对路径（用于回答时标注来源）
- page      PDF 页码（0 基，展示时 +1）
- loader    加载器类型
- is_markdown / has_table  供 splitter 选择分割策略
"""
import re
from pathlib import Path
from typing import List

from langchain_core.documents import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def load_documents(folder_path: str) -> List[Document]:
    """递归扫描 folder_path 下所有受支持的文档并加载。

    Args:
        folder_path: 文档目录（可嵌套子目录）。

    Returns:
        List[Document]：原始文档列表（尚未分割）。
    """
    folder = Path(folder_path)
    if not folder.exists():
        raise FileNotFoundError(f"目录不存在: {folder_path}")

    docs: List[Document] = []
    for path in sorted(folder.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        docs.extend(_load_single(path))

    if not docs:
        print(f"[loader] 警告：{folder_path} 中没有找到受支持的文档"
              f"（支持格式：{', '.join(sorted(SUPPORTED_EXTENSIONS))}）")
    else:
        print(f"[loader] 成功加载 {len(docs)} 个文档块（来自 {folder_path}）")
    return docs


def load_files(paths) -> List[Document]:
    """加载任意文件列表（供 Streamlit 上传文件使用）。"""
    docs: List[Document] = []
    for p in paths:
        path = Path(p)
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            docs.extend(_load_single(path))
    return docs


# ---------------------------------------------------------------------------
# 单文件加载
# ---------------------------------------------------------------------------

def _load_single(path: Path) -> List[Document]:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            return _load_pdf(path)
        if ext == ".docx":
            return [_load_docx(path)]
        if ext == ".txt":
            return [_load_text(path)]
        if ext in (".md", ".markdown"):
            return [_load_markdown(path)]
    except Exception as exc:  # 单个文件失败不影响整体
        print(f"[loader] 跳过 {path.name}（{exc}）")
        return []
    return []


def _load_pdf(path: Path) -> List[Document]:
    """按页加载 PDF，保留页码元数据。"""
    from langchain_community.document_loaders import PyPDFLoader

    loader = PyPDFLoader(str(path))
    pages = loader.load()
    docs: List[Document] = []
    for page_no, page in enumerate(pages):
        text = (page.page_content or "").strip()
        if not text:
            continue
        docs.append(Document(
            page_content=text,
            metadata={"source": str(path), "page": page_no, "loader": "pdf"},
        ))
    return docs


def _load_docx(path: Path) -> Document:
    """加载 .docx：按文档顺序抽取段落与表格，表格转为 Markdown 表格保留结构。"""
    from docx import Document as DocxDocument
    from docx.document import Document as _DocxDoc
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(str(path))

    def iter_blocks(parent):
        parent_elm = parent.element.body if isinstance(parent, _DocxDoc) else parent._element
        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    parts: List[str] = []
    has_table = False
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                parts.append(text)
        else:  # 表格
            has_table = True
            parts.append(_table_to_markdown(block))

    content = "\n\n".join(parts)
    return Document(
        page_content=content,
        metadata={
            "source": str(path),
            "loader": "docx",
            "has_table": has_table,
            "is_markdown": True,  # 表格已转为 Markdown，交给 MarkdownHeaderTextSplitter
        },
    )


def _table_to_markdown(table) -> str:
    """python-docx Table → Markdown 表格文本。"""
    rows = []
    for row in table.rows:
        cells = [(c.text or "").replace("\n", " ").strip() for c in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
    if len(rows) >= 2:  # 首行视为表头
        n = rows[0].count("|") - 1
        rows.insert(1, "|" + "---|" * n)
    return "\n".join(rows)


def _load_text(path: Path) -> Document:
    """加载 .txt：自动探测编码（utf-8 → gbk → gb18030），识别 Markdown/表格。"""
    text = _read_text(path)
    meta = {"source": str(path), "loader": "txt"}
    if _looks_like_markdown(text):
        meta["is_markdown"] = True
    if _contains_table(text):
        meta["has_table"] = True
    return Document(page_content=text, metadata=meta)


def _load_markdown(path: Path) -> Document:
    text = _read_text(path)
    meta = {
        "source": str(path),
        "loader": "md",
        "is_markdown": True,
        "has_table": _contains_table(text),
    }
    return Document(page_content=text, metadata=meta)


# ---------------------------------------------------------------------------
# 文本辅助
# ---------------------------------------------------------------------------

def _read_text(path: Path) -> str:
    for enc in ("utf-8", "gbk", "gb18030", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _contains_table(text: str) -> bool:
    """粗略检测 Markdown 表格：存在 2 行以上以 | 开头的连续/非连续行。"""
    lines = [ln for ln in text.splitlines() if ln.strip().startswith("|")]
    return len(lines) >= 2


def _looks_like_markdown(text: str) -> bool:
    """检测文本中是否含有 Markdown 标题。"""
    return bool(re.search(r"^#{1,6}\s+\S", text, re.M))
