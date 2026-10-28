"""生成 PDF / DOCX 格式的测试文档（覆盖四种文件格式的加载路径）。

用法：
    python test_samples/generate_samples.py

依赖：reportlab（PDF）、python-docx（DOCX，一般已随项目安装）
运行后会在 test_samples/ 下生成：
    《民法典_合同编_示例.docx》（含段落 + 表格）
    《民法典_合同编_示例.pdf》  （按页排版，供页码溯源演示）
"""
from pathlib import Path

HERE = Path(__file__).resolve().parent
SRC_MD = HERE / "民法典_合同编.md"


def read_source() -> str:
    return SRC_MD.read_text(encoding="utf-8")


def generate_docx() -> None:
    """生成带表格的 Word 文档：验证 docx 表格 → Markdown 表格 → MarkdownHeaderTextSplitter 路径。"""
    try:
        from docx import Document
    except ImportError:
        print("缺少 python-docx，跳过 DOCX 生成")
        return

    doc = Document()
    doc.add_heading("中华人民共和国民法典 · 合同编（节选）", level=0)

    current_section = ""
    table_rows = []
    in_table = False
    for line in read_source().splitlines():
        line = line.strip()
        if line.startswith("## "):
            current_section = line[3:].strip()
            doc.add_heading(current_section, level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("|"):
            if not in_table:
                in_table = True
                table_rows = []
            cells = [c.strip() for c in line.strip("|").split("|")]
            if cells and not all(set(c) <= {"-"} for c in cells):
                table_rows.append(cells)
        elif line.startswith(">"):
            doc.add_paragraph(line[1:].strip(), style="Intense Quote")
        elif line:
            if in_table:
                in_table = False
                if table_rows:
                    t = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                    for ri, row in enumerate(table_rows):
                        for ci, val in enumerate(row):
                            t.cell(ri, ci).text = val
                    table_rows = []
            doc.add_paragraph(line)

    if in_table and table_rows:  # 收尾处理
        t = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for ri, row in enumerate(table_rows):
            for ci, val in enumerate(row):
                t.cell(ri, ci).text = val

    out = HERE / "民法典_合同编_示例.docx"
    doc.save(out)
    print(f"[OK] 已生成 {out.name}")


def generate_pdf() -> None:
    """生成多页 PDF：验证按页加载 + 页码溯源。"""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        print("缺少 reportlab，跳过 PDF 生成（pip install reportlab）")
        return

    # 注册中文字体（reportlab 内置 CID 字体，无需外部字体文件）
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    FONT = "STSong-Light"

    styles = getSampleStyleSheet()
    for s in styles.byName.values():
        s.fontName = FONT
    title_style = ParagraphStyle("CNTitle", parent=styles["Title"], fontSize=16, leading=20, fontName=FONT)
    h1_style = ParagraphStyle("CNH1", parent=styles["Heading1"], fontSize=14, leading=18, fontName=FONT)
    h2_style = ParagraphStyle("CNH2", parent=styles["Heading2"], fontSize=12, leading=16, fontName=FONT)
    body_style = ParagraphStyle(
        "CNBody", parent=styles["BodyText"], fontSize=10.5, leading=16,
        wordWrap="CJK", spaceAfter=6,
    )

    out = HERE / "民法典_合同编_示例.pdf"
    doc = SimpleDocTemplate(str(out), pagesize=A4,
                            leftMargin=2.5 * cm, rightMargin=2.5 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)

    story = []
    story.append(Paragraph("中华人民共和国民法典 · 合同编（节选）", title_style))
    story.append(Spacer(1, 12))
    story.append(Paragraph("以下为法律条文节选，用于 RAG 检索与问答演示。", body_style))
    story.append(Spacer(1, 6))

    page_lines = 0
    for line in read_source().splitlines():
        line = line.strip()
        if not line or line.startswith(">"):
            continue
        if line.startswith("### "):
            story.append(Paragraph(line[4:], h2_style))
        elif line.startswith("## "):
            story.append(Spacer(1, 8))
            story.append(Paragraph(line[3:], h1_style))
        elif line.startswith("# "):
            continue
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if cells and all(set(c) <= {"-"} for c in cells):
                continue
            tbl = Table([cells], colWidths=[3.2 * cm] * len(cells))
            tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, "#999999"),
                                     ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                                     ("FONTNAME", (0, 0), (-1, -1), FONT)]))
            story.append(tbl)
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph(line, body_style))
        page_lines += 1
        if page_lines % 28 == 0:  # 简单分页，保证多页效果
            story.append(PageBreak())

    doc.build(story)
    print(f"[OK] 已生成 {out.name}（多页，含页码溯源信息）")


if __name__ == "__main__":
    generate_docx()
    generate_pdf()
